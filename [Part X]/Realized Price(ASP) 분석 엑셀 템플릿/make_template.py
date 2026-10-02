"""
Realized Price(ASP) 전년 대비 분석 엑셀 템플릿 생성기

세일즈 RAW 데이터를 붙여넣으면 전체 / 업체별 ASP(매출액 ÷ 수량)를 전년과 비교해 주는
엑셀 파일을 만듭니다. 필요한 열은 '열 위치'가 아니라 '헤더명'으로 찾기 때문에,
추출 데이터에 다른 열이 아무리 많아도 그대로 붙여넣으면 됩니다.

실행:  python make_template.py            ->  RealizedPrice_ASP_템플릿.xlsx 생성
필요:  pip install openpyxl pandas numpy

시트 구성
  사용방법    : 사용 순서, 계산 기준, 주의사항
  업체별 ASP  : 전체 + 업체별 ASP 전년 대비 (결과 화면)
  설정        : 열 매핑(헤더명), 비교 기간, 필터, 정렬, 데이터 점검
  RAW         : 세일즈 데이터 붙여넣는 곳 (예시 데이터 포함)
  _calc       : 내부 계산용 (숨김)
"""

import datetime as dt
import re
import shutil
import zipfile
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.formatting.rule import Rule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.styles.differential import DifferentialStyle
from openpyxl.styles.numbers import NumberFormat
from openpyxl.utils.indexed_list import IndexedList
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.formula import ArrayFormula

OUTPUT_NAME = "RealizedPrice_ASP_템플릿.xlsx"

CAPACITY = 3000                 # 업체별 표에 표시할 수 있는 최대 업체 수
SAMPLE_MARK = "※예시데이터"      # 예시 데이터 각 행의 마지막 열에 들어가는 표식
SAMPLE_MARK_COL = "S"           # 예시 데이터 표식 열 (RAW 시트)
NO_FILTER = "<>§NOFILTER§"      # 필터 미사용 시 SUMIFS 조건 (모든 행 일치)
NUM_ONLY = '">-1E+300"'         # 합계 열 자신에 거는 조건: 숫자 셀만 합산 (오류·텍스트 셀 제외)
UNUSED_COL = 16384              # 헤더를 못 찾았을 때 가리키는 빈 열 (XFD)

SH_GUIDE, SH_OUT, SH_SET, SH_RAW, SH_CALC = "사용방법", "업체별 ASP", "설정", "RAW", "_calc"
Q_OUT, Q_SET, Q_CALC = "'업체별 ASP'", "'설정'", "'_calc'"

SORT_OPTIONS = [
    "당해 매출액 큰 순",
    "ASP 증감률 낮은 순 (하락 먼저)",
    "ASP 증감률 높은 순 (상승 먼저)",
]

# 설정 시트 입력값 기본값 (예시 데이터 기준)
DEFAULT_SETTINGS = {
    "hdr_row": 1,
    "map_period": "청구일", "map_cust": "고객명", "map_qty": "판매수량", "map_amt": "순매출액",
    "cy_from": dt.datetime(2026, 1, 1), "cy_to": dt.datetime(2026, 9, 30),
    "py_from": None, "py_to": None,
    "f1_hdr": "사업부", "f1_val": None,
    "f2_hdr": "제품군", "f2_val": None,
    "f3_hdr": "통화", "f3_val": None,
    "sort": SORT_OPTIONS[0],
}

# ---------------------------------------------------------------------------
# 스타일
# ---------------------------------------------------------------------------
FONT_NAME = "맑은 고딕"
NAVY, BLUE_LIGHT, GREY_TXT = "1F3864", "D9E1F2", "595959"
F_BASE = Font(name=FONT_NAME, size=10)
F_BOLD = Font(name=FONT_NAME, size=10, bold=True)
F_TITLE = Font(name=FONT_NAME, size=15, bold=True, color=NAVY)
F_SECTION = Font(name=FONT_NAME, size=11, bold=True, color=NAVY)
F_NOTE = Font(name=FONT_NAME, size=9, color=GREY_TXT)
F_HEAD = Font(name=FONT_NAME, size=10, bold=True, color="FFFFFF")
F_INPUT = Font(name=FONT_NAME, size=10, color="0000FF")
F_CALC = Font(name=FONT_NAME, size=10, color="404040")

FILL_HEAD = PatternFill("solid", fgColor=NAVY)
FILL_SUB = PatternFill("solid", fgColor=BLUE_LIGHT)
FILL_INPUT = PatternFill("solid", fgColor="FFF2CC")
FILL_AUTO = PatternFill("solid", fgColor="F2F2F2")
FILL_TOTAL = PatternFill("solid", fgColor="EDF2F9")

THIN = Side(style="thin", color="BFBFBF")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
MED_NAVY = Side(style="medium", color=NAVY)

CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
LEFT = Alignment(horizontal="left", vertical="center")
LEFT_WRAP = Alignment(horizontal="left", vertical="center", wrap_text=True)

FMT_AMT = '#,##0;-#,##0;"-"'
FMT_ASP = '#,##0.00;-#,##0.00;"-"'
FMT_PCT = '0.0%;-0.0%;"-"'
FMT_DASP = '[Red]"▲"#,##0.00;[Blue]"▼"#,##0.00;"-"'
FMT_DPCT = '[Red]"▲"0.0%;[Blue]"▼"0.0%;"-"'
FMT_PERIOD = "yyyy-mm-dd"


def style(cell, font=F_BASE, fill=None, fmt=None, align=None, border=None):
    cell.font = font
    if fill is not None:
        cell.fill = fill
    if fmt is not None:
        cell.number_format = fmt
    if align is not None:
        cell.alignment = align
    if border is not None:
        cell.border = border
    return cell


_CELL_LIKE_NAME = re.compile(r"^[A-Za-z]{1,3}\d+$|^[RrCc]\d*$|^[Rr]\d*[Cc]\d*$")


def add_name(wb, name, ref):
    # 'dF1'처럼 셀 주소(DF1)로 읽히는 이름은 엑셀에서 쓸 수 없음
    assert not _CELL_LIKE_NAME.match(name), f"이름이 셀 주소처럼 보입니다: {name}"
    wb.defined_names[name] = DefinedName(name, attr_text=ref)


def font_rule(formula, color, bold=False):
    return Rule(type="expression", formula=[formula],
                dxf=DifferentialStyle(font=Font(color=color, bold=bold)))


# ---------------------------------------------------------------------------
# 예시 데이터 (가상의 B2B 소재 회사 세일즈 추출 데이터)
# ---------------------------------------------------------------------------
SAMPLE_HEADERS = ["회사코드", "사업부", "판매조직", "고객코드", "고객명", "제품군", "제품코드", "제품명",
                  "청구일", "청구연월", "판매수량", "단위", "총매출액", "할인액", "리베이트", "순매출액",
                  "통화", "전표번호", SAMPLE_MARK]

PRODUCTS = [  # 제품군, 제품코드, 제품명, 전년 리스트가(원/KG)
    ("PP", "M-1101", "PP 범용사출 H110", 1650),
    ("PP", "M-1102", "PP 고강성 H220", 1820),
    ("PP", "M-1103", "PP 투명 R330", 1940),
    ("PE", "M-1201", "HDPE 블로우 B500", 1580),
    ("PE", "M-1202", "LLDPE 필름 F120", 1710),
    ("ABS", "M-2101", "ABS 범용 G300", 2650),
    ("ABS", "M-2102", "ABS 내열 HR500", 2980),
    ("PC", "M-2201", "PC 일반 C100", 3900),
    ("PC", "M-2202", "PC 난연 FR200", 4350),
]
LIST_CHANGE = {"PP": 0.030, "PE": -0.015, "ABS": 0.040, "PC": 0.020}   # 당해 리스트가 변동률
DIVISION = {"PP": "범용수지사업부", "PE": "범용수지사업부", "ABS": "EP사업부", "PC": "EP사업부"}

CUSTOMERS = [
    # 고객명, 판매조직, 제품군, 월 물량(KG), 전년 할인율, 당해 할인율 변화(%p), 리베이트율,
    # 당해 물량 배수, 당해 제품군별 물량 배수(믹스 변화), 거래 시작월, 거래 종료월
    ("가온소재(주)", "국내영업1팀", ["PP", "PE"], 18000, 0.06, -0.010, 0.00, 1.05, {}, None, None),
    ("나래산업(주)", "국내영업1팀", ["PP"], 22000, 0.08, 0.025, 0.01, 1.10, {}, None, None),
    ("다솜케미칼", "국내영업2팀", ["ABS", "PC"], 9000, 0.05, 0.000, 0.00, 0.95, {}, None, None),
    ("라온테크(주)", "국내영업2팀", ["ABS"], 12000, 0.07, 0.040, 0.00, 1.15, {}, None, None),
    ("마루정밀", "국내영업1팀", ["PC"], 5000, 0.04, -0.005, 0.00, 1.00, {}, None, None),
    ("바른물산", "국내영업2팀", ["PE"], 15000, 0.09, 0.010, 0.02, 0.90, {}, None, None),
    ("새봄플라스틱", "국내영업1팀", ["PP", "ABS"], 11000, 0.06, 0.000, 0.00, 1.00, {"PP": 0.7, "ABS": 1.6}, None, None),
    ("아름화학(주)", "국내영업2팀", ["PE", "PP"], 20000, 0.10, 0.005, 0.01, 1.02, {}, None, None),
    ("자연엠텍", "국내영업1팀", ["PC", "ABS"], 6000, 0.05, -0.015, 0.00, 1.08, {}, None, None),
    ("차오름산업", "국내영업2팀", ["PP"], 8000, 0.07, 0.000, 0.00, 1.00, {}, None, "2025-08"),
    ("큰별코리아", "해외영업팀", ["PP", "PE"], 25000, 0.11, 0.015, 0.02, 1.12, {}, None, None),
    ("토담정공(주)", "국내영업1팀", ["ABS"], 4000, 0.03, 0.000, 0.00, 0.97, {}, None, None),
    ("파란하이텍", "국내영업2팀", ["PC"], 7000, 0.06, 0.030, 0.00, 1.20, {}, None, None),
    ("하늘포장(주)", "국내영업1팀", ["PE"], 16000, 0.08, -0.010, 0.00, 1.00, {}, None, None),
    ("한울몰드", "국내영업2팀", ["PP", "ABS"], 9000, 0.06, 0.005, 0.00, 1.03, {}, None, None),
    ("온새미사출", "국내영업1팀", ["PP"], 7000, 0.05, 0.000, 0.00, 0.92, {}, None, None),
    ("누리오토파츠", "국내영업2팀", ["ABS", "PC"], 10000, 0.07, 0.020, 0.01, 1.00, {"PC": 0.6, "ABS": 1.4}, None, None),
    ("든솔케이스", "국내영업1팀", ["ABS"], 3000, 0.04, 0.000, 0.00, 1.00, {}, "2025-10", None),
    ("미르패키징", "국내영업2팀", ["PE", "PP"], 12000, 0.09, -0.005, 0.00, 1.06, {}, None, None),
    ("별하전자(주)", "국내영업1팀", ["PC"], 2500, 0.03, 0.000, 0.00, 1.10, {}, None, None),
    ("Sinar Plastik (ID)", "해외영업팀", ["PP"], 30000, 0.12, 0.020, 0.00, 1.05, {}, None, None),
    ("Hanoi Molding (VN)", "해외영업팀", ["ABS", "PP"], 14000, 0.10, 0.000, 0.00, 1.00, {}, "2026-02", None),
]


def make_sample_data(seed=7):
    """2025-01 ~ 2026-09 청구 기준 가상 세일즈 데이터 (반품, 단가 소급정산 포함)."""
    rng = np.random.default_rng(seed)
    months = pd.period_range("2025-01", "2026-09", freq="M")
    rows = []
    for ci, (name, team, groups, base_qty, disc, d_disc, rebate, growth, mix, start, end) in enumerate(CUSTOMERS):
        code = f"C{10001 + ci * 7}"
        # 고객별로 제품군마다 1~2개 품목 구매
        items = []
        for g in groups:
            cands = [p for p in PRODUCTS if p[0] == g]
            k = 1 if len(cands) == 1 else int(rng.integers(1, 3))
            idx = rng.choice(len(cands), size=k, replace=False)
            items += [cands[i] for i in sorted(idx)]
        for m in months:
            if start and m < pd.Period(start, "M"):
                continue
            if end and m > pd.Period(end, "M"):
                continue
            cy = m.year == 2026
            for grp, pcode, pname, list_py in items:
                if rng.random() > 0.88:          # 주문 없는 달
                    continue
                qty_month = base_qty / len(items) * (1 + 0.12 * np.sin(m.month / 12 * 2 * np.pi))
                qty_month *= rng.lognormal(0, 0.22)
                if cy:
                    qty_month *= growth * mix.get(grp, 1.0)
                list_price = list_py * (1 + LIST_CHANGE[grp]) if cy else list_py
                list_price *= 1 + 0.006 * np.sin(m.month)          # 월별 소폭 변동
                d = disc + (d_disc if cy else 0.0)
                n_inv = 2 if rng.random() < 0.3 else 1     # 한 달에 전표 1~2건
                for _ in range(n_inv):
                    qty = int(round(qty_month / n_inv / 25.0) * 25)
                    if qty <= 0:
                        continue
                    day = int(rng.integers(1, 29))
                    rows.append(_sale_row(code, name, team, grp, pcode, pname, m, day, qty, list_price, d, rebate))
                    if rng.random() < 0.015:     # 반품: 다음 달 일부 수량 회수
                        nm = m + 1
                        if nm <= months[-1]:
                            rq = -int(round(qty * rng.uniform(0.05, 0.2) / 25.0) * 25 or 25)
                            rows.append(_sale_row(code, name, team, grp, pcode, pname, nm, int(rng.integers(1, 29)),
                                                  rq, list_price, d, rebate))
    # 단가 소급 정산 (수량 0, 금액만 차감) - 나래산업(주)
    for mon, amt in (("2026-04", 4_200_000), ("2026-07", 3_600_000)):
        rows.append({"회사코드": 1000, "사업부": DIVISION["PP"], "판매조직": "국내영업1팀",
                     "고객코드": "C10008", "고객명": "나래산업(주)", "제품군": "PP", "제품코드": "M-1101",
                     "제품명": "PP 범용사출 H110", "청구일": pd.Period(mon, "M").to_timestamp() + pd.Timedelta(days=27),
                     "판매수량": 0, "단위": "KG", "총매출액": 0, "할인액": 0, "리베이트": amt,
                     "순매출액": -amt, "통화": "KRW"})
    df = pd.DataFrame(rows)
    df["청구일"] = pd.to_datetime(df["청구일"])
    df["청구연월"] = df["청구일"].dt.year * 100 + df["청구일"].dt.month
    df = df.sort_values(["청구일", "고객코드", "제품코드"], kind="stable").reset_index(drop=True)
    df["전표번호"] = [f"INV{d:%y%m}-{i + 1:05d}" for i, d in enumerate(df["청구일"])]
    df[SAMPLE_MARK] = SAMPLE_MARK
    return df[SAMPLE_HEADERS]


def _sale_row(code, name, team, grp, pcode, pname, month, day, qty, list_price, disc, rebate):
    gross = round(qty * list_price)
    discount = round(gross * disc)
    reb = round(gross * rebate)
    return {"회사코드": 1000, "사업부": DIVISION[grp], "판매조직": team, "고객코드": code, "고객명": name,
            "제품군": grp, "제품코드": pcode, "제품명": pname,
            "청구일": month.to_timestamp() + pd.Timedelta(days=day - 1),
            "판매수량": qty, "단위": "KG", "총매출액": gross, "할인액": discount, "리베이트": reb,
            "순매출액": gross - discount - reb, "통화": "KRW"}


# ---------------------------------------------------------------------------
# 시트 작성
# ---------------------------------------------------------------------------
def build_guide(ws):
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 2
    ws.column_dimensions["B"].width = 4
    ws.column_dimensions["C"].width = 26
    ws.column_dimensions["D"].width = 92
    style(ws["B1"], F_TITLE).value = "Realized Price (ASP) 전년 대비 분석 — 사용 방법"
    style(ws["B2"], F_NOTE).value = "세일즈 데이터를 붙여넣으면 전체·업체별 ASP(매출액 ÷ 수량)를 전년과 비교합니다."

    r = 4
    style(ws.cell(r, 2), F_SECTION).value = "시작하기 — 3단계"
    steps = [
        ("1", "RAW 시트에 붙여넣기",
         "RAW 시트에서 Ctrl+A → Delete로 예시 데이터를 지운 뒤, 추출한 세일즈 데이터를 A1 셀에 "
         "헤더(열 이름) 포함 그대로 붙여넣습니다. 열이 많거나 순서가 달라도 괜찮습니다 — "
         "필요한 열은 '헤더명'으로 찾아서 씁니다."),
        ("2", "설정 시트 노란 칸 입력",
         "② 열 매핑: 기간·업체·수량·매출액 헤더를 드롭다운에서 선택  /  ③ 비교 기간: 당해 시작·종료 입력 "
         "(전년은 자동으로 1년 전)  /  ④ 필터: 사업부·제품군·통화 등으로 좁힐 때만  /  "
         "⑥ 데이터 점검이 모두 ✔ 인지 확인"),
        ("3", "업체별 ASP 시트에서 확인",
         "맨 위 '전체' 행이 전체 ASP 전년 대비, 그 아래가 업체별입니다. 정렬 기준(매출액 / ASP 증감률)은 "
         "설정 ⑤에서 바꿀 수 있습니다."),
    ]
    for no, title, desc in steps:
        r += 1
        style(ws.cell(r, 2), F_HEAD, FILL_HEAD, align=CENTER).value = no
        style(ws.cell(r, 3), F_BOLD, align=LEFT_WRAP).value = title
        style(ws.cell(r, 4), F_BASE, align=LEFT_WRAP).value = desc
        ws.row_dimensions[r].height = 46

    r += 2
    style(ws.cell(r, 2), F_SECTION).value = "계산 기준"
    defs = [
        ("ASP", "매출액 합계 ÷ 수량 합계  (기간·업체별로 먼저 합산한 뒤 나눔 = 수량 가중 평균 단가)"),
        ("ASP 증감 / 증감률", "당해 ASP − 전년 ASP   /   당해 ASP ÷ 전년 ASP − 1"),
        ("당해 비중", "업체 당해 매출액 ÷ 전체 당해 매출액"),
        ("기간 판정", "기간 열 값이 [시작, 종료] 안에 드는 행만 합산 (날짜 · 연도 2026 · 연월 202609 숫자 모두 가능)"),
        ("비고", "신규(전년 수량 없음) / 당해 수량 없음 → 해당 업체는 ASP 증감을 계산하지 않음"),
    ]
    for k, v in defs:
        r += 1
        style(ws.cell(r, 3), F_BOLD, align=LEFT_WRAP).value = k
        style(ws.cell(r, 4), F_BASE, align=LEFT_WRAP).value = v
        ws.row_dimensions[r].height = 20

    r += 2
    style(ws.cell(r, 2), F_SECTION).value = "주의할 점"
    notes = [
        "수량 단위(KG / EA 등)나 통화가 섞여 있으면 ASP가 왜곡됩니다 → 설정 ④ 필터로 하나만 남기세요.",
        "업체 ASP 증감에는 '같은 제품의 가격 변화'와 '제품 믹스 변화'가 함께 들어 있습니다. "
        "제품군별로 보려면 필터에 제품군을 지정하세요.",
        "반품(− 수량)과 단가 소급 정산(수량 0, 금액만 있는 행)도 그대로 합산됩니다 = 실제 실현 단가.",
        "날짜·숫자가 텍스트로 붙여넣어지면 합계에서 빠집니다 → 설정 ⑥ 데이터 점검에서 확인하세요.",
        "업체 목록 자동 생성에는 Excel 2021 또는 Microsoft 365가 필요합니다(UNIQUE 함수). "
        "이전 버전은 설정 ⑥의 안내대로 업체 목록을 직접 붙여넣으세요.",
        f"업체는 최대 {CAPACITY:,}개까지 표시됩니다. 수식이 들어 있는 칸(흰색·회색)은 지우지 마세요.",
    ]
    for n in notes:
        r += 1
        ws.merge_cells(start_row=r, start_column=3, end_row=r, end_column=4)
        style(ws.cell(r, 3), F_BASE, align=LEFT_WRAP).value = "•  " + n
        ws.row_dimensions[r].height = 30

    r += 2
    style(ws.cell(r, 2), F_SECTION).value = "표시 안내"
    legend = [
        (FILL_INPUT, F_INPUT, "입력", "노란 칸 = 직접 입력하거나 드롭다운에서 고르는 칸"),
        (None, Font(name=FONT_NAME, size=10, bold=True, color="FF0000"), "▲ 1.2%", "전년 대비 상승"),
        (None, Font(name=FONT_NAME, size=10, bold=True, color="0000FF"), "▼ 1.2%", "전년 대비 하락"),
        (FILL_AUTO, F_CALC, "자동", "회색 칸 = 자동 계산 (수정 불필요)"),
    ]
    for fill, font, sample, desc in legend:
        r += 1
        style(ws.cell(r, 3), font, fill, align=CENTER, border=BOX).value = sample
        style(ws.cell(r, 4), F_BASE, align=LEFT).value = desc


def build_settings(ws, settings):
    ws.sheet_view.showGridLines = False
    widths = {"A": 2, "B": 22, "C": 26, "D": 24, "E": 24, "F": 70, "G": 2, "H": 12, "I": 18}
    for col, w in widths.items():
        ws.column_dimensions[col].width = w
    ws.column_dimensions["H"].hidden = True    # 내부 계산값
    ws.column_dimensions["I"].hidden = True

    style(ws["B1"], F_TITLE).value = "설정"
    style(ws["B2"], F_NOTE).value = "노란 칸만 입력 / 선택하세요. 나머지는 자동으로 계산됩니다."
    style(ws["H1"], F_NOTE).value = "내부 계산용"

    def section(row, text):
        style(ws.cell(row, 2), F_SECTION).value = text

    def header(row, labels):
        for i, text in enumerate(labels):
            style(ws.cell(row, 2 + i), F_HEAD, FILL_HEAD, align=CENTER, border=BOX).value = text

    def inp(ref, value, fmt=None):
        c = style(ws[ref], F_INPUT, FILL_INPUT, fmt=fmt, align=LEFT, border=BOX)
        c.value = value
        return c

    def calc(ref, formula, fmt=None, align=LEFT, font=F_CALC):
        c = style(ws[ref], font, FILL_AUTO, fmt=fmt, align=align, border=BOX)
        c.value = formula
        return c

    def desc(ref, text):
        style(ws[ref], F_NOTE, align=LEFT_WRAP).value = text

    # ① RAW 데이터 구조 ------------------------------------------------------
    section(4, "① RAW 데이터 구조")
    style(ws["B5"], F_BOLD, align=LEFT, border=BOX).value = "헤더 행 번호"
    inp("C5", settings["hdr_row"], "0")
    calc("E5", '=IF(COUNTA(hdrRow)=0,"✖ 헤더 행이 비어 있음","✔ 헤더 "&COUNTA(hdrRow)&"개")')
    desc("F5", "RAW 시트에서 열 이름(헤더)이 있는 행 번호. 보통 1 (리포트 제목 줄이 위에 있으면 그 아래 행 번호)")
    ws["H5"] = "=IFERROR(MIN(MAX(1,INT(C5)),1000),1)"

    # ② 열 매핑 --------------------------------------------------------------
    section(7, "② 열 매핑 — RAW 헤더명 선택")
    header(8, ["항목", "RAW 헤더명 (드롭다운)", "찾은 위치", "상태", "설명"])
    mapping = [
        (9, "기간", "map_period", "기간 판정에 쓰는 열: 날짜 / 연도(2026) / 연월(202609) 형식의 숫자"),
        (10, "업체", "map_cust", "업체별로 묶는 기준 열 (고객명, 고객코드 등)"),
        (11, "수량", "map_qty", "ASP의 분모 (판매수량)"),
        (12, "매출액", "map_amt", "ASP의 분자 — Realized price 기준 금액 (예: 할인·리베이트 차감 후 순매출액)"),
    ]
    for row, label, key, text in mapping:
        style(ws.cell(row, 2), F_BOLD, align=LEFT, border=BOX).value = label
        inp(f"C{row}", settings[key])
        calc(f"D{row}", f'=IF($H{row}={UNUSED_COL},"",SUBSTITUTE(ADDRESS(1,$H{row},4),"1","")&" 열")',
             align=CENTER)
        calc(f"E{row}", f'=IF($C{row}="","✖ 미선택",IF($H{row}={UNUSED_COL},"✖ 헤더 없음",'
                        f'IF(COUNTIF(hdrRow,$C{row})>1,"⚠ 같은 헤더 여러 개(첫 열 사용)","✔ 정상")))')
        desc(f"F{row}", text)
        ws[f"H{row}"] = f'=IF($C{row}="",{UNUSED_COL},IFERROR(MATCH($C{row},hdrRow,0),{UNUSED_COL}))'

    # ③ 비교 기간 ------------------------------------------------------------
    section(14, "③ 비교 기간")
    header(15, ["구분", "시작", "종료", "표시", "설명"])
    style(ws["B16"], F_BOLD, align=LEFT, border=BOX).value = "당해 (CY)"
    inp("C16", settings["cy_from"], FMT_PERIOD)
    inp("D16", settings["cy_to"], FMT_PERIOD)
    desc("F16", "기간 열과 같은 형식으로 입력: 날짜 2026-01-01  /  연도 2026  /  연월 202601")
    style(ws["B17"], F_BOLD, align=LEFT, border=BOX).value = "전년 직접 지정 (선택)"
    inp("C17", settings["py_from"], FMT_PERIOD)
    inp("D17", settings["py_to"], FMT_PERIOD)
    desc("F17", "비워 두면 당해의 1년 전 기간이 자동으로 적용됩니다. 다른 기간과 비교할 때만 입력")
    style(ws["B18"], F_BOLD, align=LEFT, border=BOX).value = "전년 (PY) 적용"
    for col in ("C", "D"):
        x = f"{col}16"
        auto = (f'IF(NOT(ISNUMBER({x})),"",IF({x}<3000,{x}-1,IF({x}<100000,EDATE({x},-12),'
                f'IF({x}<1000000,{x}-100,{x}-10000))))')
        calc(f"{col}18", f'=IF({col}17<>"",{col}17,{auto})', FMT_PERIOD)
    desc("F18", "실제 비교에 쓰이는 전년 기간 (자동 계산)")
    for row in (16, 18):
        fmt_a = _period_text(f"C{row}")
        fmt_b = _period_text(f"D{row}")
        calc(f"E{row}", f'=IF(OR(NOT(ISNUMBER(C{row})),NOT(ISNUMBER(D{row}))),"",'
                        f'{fmt_a}&IF(C{row}=D{row},""," ~ "&{fmt_b}))', align=CENTER)
    # 연도·연월처럼 날짜가 아닌 숫자는 날짜 서식 대신 숫자로 표시
    ws.conditional_formatting.add("C16:D18", Rule(
        type="expression", formula=["AND(ISNUMBER(C16),OR(C16<3000,C16>=100000))"],
        dxf=DifferentialStyle(numFmt=NumberFormat(numFmtId=200, formatCode="0"))))

    # ④ 필터 ----------------------------------------------------------------
    section(20, "④ 필터 (선택 — 비우면 전체 데이터)")
    header(21, ["구분", "RAW 헤더명 (드롭다운)", "조건 값", "상태", "설명"])
    filter_desc = [
        "예) 사업부 = 범용수지사업부   /   조건 값에 <>내부거래 처럼 쓰면 '제외'",
        "예) 제품군 = PP   /   PP* 처럼 * 를 쓰면 'PP로 시작하는 값'",
        "예) 통화 = KRW  또는  단위 = KG  (단위·통화가 섞여 있으면 꼭 지정)",
    ]
    for i, row in enumerate((22, 23, 24)):
        n = i + 1
        style(ws.cell(row, 2), F_BOLD, align=LEFT, border=BOX).value = f"필터 {n}"
        inp(f"C{row}", settings[f"f{n}_hdr"])
        inp(f"D{row}", settings[f"f{n}_val"])
        calc(f"E{row}", f'=IF(AND($C{row}="",$D{row}=""),"– 미사용",IF($D{row}="","– 미적용 (조건 값 없음)",'
                        f'IF($C{row}="","⚠ 헤더 미선택",IF(ISNA(MATCH($C{row},hdrRow,0)),"✖ 헤더 없음",'
                        f'"✔ 적용"))))')
        desc(f"F{row}", filter_desc[i])
        ws[f"H{row}"] = (f'=IF(OR($C{row}="",$D{row}=""),{UNUSED_COL},'
                         f'IFERROR(MATCH($C{row},hdrRow,0),{UNUSED_COL}))')
        ws[f"I{row}"] = f'=IF($H{row}={UNUSED_COL},"{NO_FILTER}",$D{row})'
    style(ws["B25"], F_BOLD, align=LEFT, border=BOX).value = "적용 중인 필터"
    parts = "&".join(f'IF(LEFT($E{r},1)="✔",$C{r}&IF(LEFT($D{r},2)="<>"," ≠ "&MID($D{r},3,999),'
                     f'" = "&$D{r})&"  ·  ","")' for r in (22, 23, 24))
    ws.merge_cells("C25:E25")
    calc("C25", f'=IF(COUNTIF($E$22:$E$24,"✔*")=0,"없음 (전체 데이터)",LEFT({parts},LEN({parts})-5))')

    # ⑤ 표시 옵션 ------------------------------------------------------------
    section(27, "⑤ 표시 옵션")
    style(ws["B28"], F_BOLD, align=LEFT, border=BOX).value = "업체 정렬 기준"
    ws.merge_cells("C28:D28")
    inp("C28", settings["sort"])
    desc("F28", "업체별 ASP 표의 정렬 순서 (드롭다운)")
    ws["H28"] = "=IFERROR(MATCH(C28,sortList,0),1)"

    # ⑥ 데이터 점검 ----------------------------------------------------------
    section(30, "⑥ 데이터 점검")
    header(31, ["점검 항목", "결과", "", "상태", "확인 / 조치 방법"])
    ws.merge_cells("C31:D31")
    period_ok = "AND(ISNUMBER(C16),ISNUMBER(D16),ISNUMBER(C18),ISNUMBER(D18))"
    flt = "dFlt1,critFlt1,dFlt2,critFlt2,dFlt3,critFlt3"
    checks = [
        ("RAW 데이터 행 수",
         "=MAX(COUNTA(dPer),COUNTA(dCust),COUNTA(dQty),COUNTA(dAmt))", '#,##0"행"',
         '=IF(C32>0,"✔ 정상","✖ 데이터 없음")',
         "RAW 시트의 헤더 행(보통 1행)부터 데이터를 헤더 포함 붙여넣으세요."),
        ("필수 열 찾기",
         '=COUNTIF($E$9:$E$12,"✔*")+COUNTIF($E$9:$E$12,"⚠*")&" / 4 열"', None,
         '=IF(COUNTIF($E$9:$E$12,"✖*")>0,"✖ 확인 필요",IF(COUNTIF($E$9:$E$12,"⚠*")>0,"⚠ 중복 헤더","✔ 정상"))',
         "② 열 매핑에서 RAW 헤더명을 드롭다운으로 다시 선택하세요."),
        ("비교 기간 입력",
         f'=IF({period_ok},"입력됨","입력 필요")', None,
         f'=IF({period_ok},IF(AND(C16<=D16,C18<=D18),"✔ 정상","✖ 시작 > 종료"),"✖ 확인 필요")',
         "③ 비교 기간을 날짜(2026-01-01) / 연도(2026) / 연월(202601) 숫자로 입력하세요."),
        ("기간 열 형식",
         '="날짜·숫자 "&TEXT(COUNT(dPer),"#,##0")&"행  /  텍스트 "&TEXT(SUMPRODUCT(--ISTEXT(dPer)),"#,##0")&"행"',
         None,
         '=IF(COUNTA(dPer)=0,"✖ 데이터 없음",IF(SUMPRODUCT(--ISTEXT(dPer))>0,"⚠ 텍스트 포함","✔ 정상"))',
         "텍스트로 된 날짜는 기간 판정에서 빠집니다 → RAW에서 그 열 선택 > 데이터 > 텍스트 나누기 > 마침"),
        ("수량·매출액 숫자 형식",
         '="텍스트 "&TEXT(SUMPRODUCT(--ISTEXT(dQty))+SUMPRODUCT(--ISTEXT(dAmt)),"#,##0")&"셀"', None,
         '=IF(SUMPRODUCT(--ISTEXT(dQty))+SUMPRODUCT(--ISTEXT(dAmt))>0,"⚠ 텍스트 포함","✔ 정상")',
         "숫자가 텍스트로 저장된 셀은 합계에서 빠집니다 → 그 열 선택 > 데이터 > 텍스트 나누기 > 마침"),
        ("오류 값 셀",
         "=SUMPRODUCT(--ISERROR(dPer))+SUMPRODUCT(--ISERROR(dCust))+SUMPRODUCT(--ISERROR(dQty))"
         "+SUMPRODUCT(--ISERROR(dAmt))", '#,##0"셀"',
         '=IF(C37>0,"⚠ 오류 셀 있음","✔ 정상")',
         "오류 값(N/A, VALUE! 등) 셀은 합계에서 빠집니다 → RAW에서 홈 > 찾기 및 선택 > 이동 옵션 > "
         "상수 > 오류만 체크 → 확인 → Delete"),
        ("당해 기간 데이터",
         f'=COUNTIFS(dPer,">="&cyFrom,dPer,"<="&cyTo,{flt})', '#,##0"행"',
         '=IF(C38>0,"✔ 정상","✖ 해당 행 없음")',
         "당해 기간·필터에 맞는 행이 없습니다 → ③ 기간 형식과 값, ④ 필터를 확인하세요."),
        ("전년 기간 데이터",
         f'=COUNTIFS(dPer,">="&pyFrom,dPer,"<="&pyTo,{flt})', '#,##0"행"',
         '=IF(C39>0,"✔ 정상","✖ 해당 행 없음")',
         "전년 기간·필터에 맞는 행이 없습니다 → 전년 데이터가 RAW에 함께 들어 있는지 확인하세요."),
        ("필터",
         '=COUNTIF($E$22:$E$24,"✔*")&"개 적용"', None,
         '=IF(COUNTIF($E$22:$E$24,"✖*")+COUNTIF($E$22:$E$24,"⚠*")>0,"⚠ 확인 필요","✔ 정상")',
         "④ 필터의 헤더명이 RAW에 있는지 확인하세요."),
        ("업체 목록 생성",
         f'=IF(ISERROR({Q_CALC}!$A$2),"목록 오류",IF(dynNA,"직접 입력 ","자동 생성 ")&TEXT(nList,"#,##0")&"개")',
         None,
         f'=IF(ISERROR({Q_CALC}!$A$2),"✖ _calc 시트 A열 확인",IF(dynNA,IF(nList=0,"✖ Excel 2021 이상 필요",'
         '"✔ 직접 입력 목록"),"✔ 정상"))',
         "Excel 2019 이하는 자동 생성이 안 됩니다 → _calc 시트 숨기기 취소 → A2부터 업체 목록 붙여넣기 "
         "(RAW의 업체 열 복사 → 데이터 > 중복된 항목 제거)"),
        ("업체 수",
         f'=TEXT(nComp,"#,##0")&"개 표시  /  한도 {CAPACITY:,}"', None,
         f'=IF(nList>{CAPACITY},"⚠ 한도 초과","✔ 정상")',
         f"업체가 {CAPACITY:,}개를 넘으면 나머지는 표시되지 않습니다 → ④ 필터로 나눠서 보세요."),
        ("합계 검증 (업체 합 = 전체)",
         '="매출 차이 "&TEXT(H43,"#,##0")&"  /  수량 차이 "&TEXT(I43,"#,##0")', None,
         '=IF(AND(ROUND(H43,0)=0,ROUND(I43,0)=0),"✔ 일치","⚠ 불일치")',
         None),
        ("예시 데이터",
         f'=IF(COUNTIF(RAW!${SAMPLE_MARK_COL}:${SAMPLE_MARK_COL},"{SAMPLE_MARK}")>0,"예시 데이터 사용 중","없음")',
         None,
         f'=IF(COUNTIF(RAW!${SAMPLE_MARK_COL}:${SAMPLE_MARK_COL},"{SAMPLE_MARK}")>0,"⚠ 예시 데이터","✔ 정상")',
         "실제 데이터를 붙여넣기 전에 RAW 시트 전체를 지우세요 (Ctrl+A → Delete)."),
    ]
    for i, (label, result, fmt, status, guide) in enumerate(checks):
        row = 32 + i
        style(ws.cell(row, 2), F_BOLD, align=LEFT, border=BOX).value = label
        ws.merge_cells(f"C{row}:D{row}")
        calc(f"C{row}", result, fmt)
        calc(f"E{row}", status)
        if guide is not None:
            desc(f"F{row}", guide)
    # 합계 검증: 업체 합계와 전체 합계의 차이, 업체 열이 빈 행의 금액
    ws["H43"] = (f"=ABS(SUM({Q_CALC}!$C$2:$C${CAPACITY + 1})-{Q_OUT}!$D$8)"
                 f"+ABS(SUM({Q_CALC}!$D$2:$D${CAPACITY + 1})-{Q_OUT}!$E$8)")
    ws["I43"] = (f"=ABS(SUM({Q_CALC}!$E$2:$E${CAPACITY + 1})-{Q_OUT}!$G$8)"
                 f"+ABS(SUM({Q_CALC}!$F$2:$F${CAPACITY + 1})-{Q_OUT}!$H$8)")
    blank_amt = (f'SUMIFS(dAmt,dCust,"",dPer,">="&cyFrom,dPer,"<="&cyTo,dAmt,{NUM_ONLY},{flt})'
                 f'+SUMIFS(dAmt,dCust,"",dPer,">="&pyFrom,dPer,"<="&pyTo,dAmt,{NUM_ONLY},{flt})')
    style(ws["F43"], F_NOTE, align=LEFT_WRAP).value = (
        f'="업체명이 빈 행의 매출액(당해+전년): "&TEXT({blank_amt},"#,##0")'
        f'&"  → 업체 열이 빈 행은 업체별 표에서 빠집니다"')

    style(ws["B45"], F_BOLD, align=LEFT, border=BOX).value = "종합"
    ws.merge_cells("C45:D45")
    calc("C45", '=COUNTIF($E$32:$E$44,"✖*")+COUNTIF($E$32:$E$44,"⚠*")', '0"건"')
    calc("E45", '=IF(C45=0,"✔ 이상 없음","⚠ "&C45&"건 확인 필요")', font=F_BOLD)

    # 상태 색상
    rng = "E5:E45"
    ws.conditional_formatting.add(rng, font_rule('LEFT(E5,1)="✔"', "00703C"))
    ws.conditional_formatting.add(rng, font_rule('LEFT(E5,1)="✖"', "C00000", bold=True))
    ws.conditional_formatting.add(rng, font_rule('LEFT(E5,1)="⚠"', "C65911", bold=True))
    ws.conditional_formatting.add(rng, font_rule('LEFT(E5,1)="–"', "808080"))

    # 드롭다운
    dv_hdr = DataValidation(type="list", formula1="hdrList", allow_blank=True, showErrorMessage=False)
    dv_hdr.add("C9:C12")
    dv_hdr.add("C22:C24")
    dv_sort = DataValidation(type="list", formula1="sortList", allow_blank=False)
    dv_sort.add("C28")
    dv_row = DataValidation(type="whole", operator="between", formula1="1", formula2="1000",
                            errorTitle="헤더 행 번호", error="1~1000 사이 숫자를 입력하세요.")
    dv_row.add("C5")
    for dv in (dv_hdr, dv_sort, dv_row):
        ws.add_data_validation(dv)
    ws.freeze_panes = "A3"


def _period_text(ref):
    """기간 값 표시: 날짜면 yyyy.mm.dd, 연도·연월이면 숫자 그대로."""
    return f'IF(AND({ref}>=3000,{ref}<100000),TEXT({ref},"yyyy.mm.dd"),TEXT({ref},"0"))'


def build_output(ws):
    ws.sheet_view.showGridLines = False
    widths = {"A": 1.5, "B": 5, "C": 28, "D": 15, "E": 15, "F": 9, "G": 13, "H": 13,
              "I": 12, "J": 12, "K": 12, "L": 10, "M": 22, "N": 8}
    for col, w in widths.items():
        ws.column_dimensions[col].width = w
    ws.column_dimensions["N"].hidden = True

    style(ws["B1"], F_TITLE).value = "업체별 Realized Price (ASP) — 전년 대비"
    style(ws["B2"], F_NOTE).value = (
        f'="비교 기간   당해 "&cyLabel&"   vs   전년 "&pyLabel&"      |      ASP = "&{Q_SET}!$C$12'
        f'&" ÷ "&{Q_SET}!$C$11&"      |      필터: "&{Q_SET}!$C$25')
    style(ws["B3"], Font(name=FONT_NAME, size=10, bold=True, color="C00000")).value = (
        f'="데이터 점검: "&{Q_SET}!$E$45&IF({Q_SET}!$C$45=0,"","   → [설정] 시트 ⑥ 데이터 점검을 확인하세요")')
    style(ws["B4"], Font(name=FONT_NAME, size=13, bold=True, color="262626")).value = (
        '=IF(OR(I8="",J8=""),"전체 ASP: 계산할 수 없습니다 (설정 시트 확인)",'
        '"전체 ASP   전년 "&TEXT(I8,"#,##0.00")&"   →   당해 "&TEXT(J8,"#,##0.00")&"     ("'
        '&IF(L8="","-",IF(L8>0,"▲ ",IF(L8<0,"▼ ",""))&TEXT(ABS(L8),"0.0%"))&")")')
    ws.row_dimensions[4].height = 24
    ws.conditional_formatting.add("B3", font_rule('LEFT($B$3,9)="데이터 점검: ✔"', "00703C", bold=True))
    ws.conditional_formatting.add("B4", font_rule("AND(ISNUMBER($L$8),$L$8>0)", "C00000", bold=True))
    ws.conditional_formatting.add("B4", font_rule("AND(ISNUMBER($L$8),$L$8<0)", "0000C0", bold=True))

    # 머리글 (2단)
    for ref, text in (("B6", "No"), ("C6", "업체"), ("M6", "비고")):
        ws.merge_cells(f"{ref}:{ref[0]}7")
        style(ws[ref], F_HEAD, FILL_HEAD, align=CENTER, border=BOX).value = text
        style(ws[f"{ref[0]}7"], F_HEAD, FILL_HEAD, border=BOX)
    for rng_, text in (("D6:F6", "매출액"), ("G6:H6", "수량"), ("I6:L6", "ASP  (매출액 ÷ 수량)")):
        ws.merge_cells(rng_)
        first, last = rng_.split(":")
        for col in range(ord(first[0]), ord(last[0]) + 1):
            style(ws[f"{chr(col)}6"], F_HEAD, FILL_HEAD, align=CENTER, border=BOX)
        ws[first].value = text
    for col, text in zip("DEFGHIJKL", ["전년", "당해", "당해 비중", "전년", "당해", "전년", "당해", "증감", "증감률"]):
        style(ws[f"{col}7"], F_BOLD, FILL_SUB, align=CENTER, border=BOX).value = text
    ws.row_dimensions[6].height = 20
    ws.row_dimensions[7].height = 20

    # 전체 행
    flt = "dFlt1,critFlt1,dFlt2,critFlt2,dFlt3,critFlt3"
    total = {
        "C": "전체",
        "D": f'=SUMIFS(dAmt,dPer,">="&pyFrom,dPer,"<="&pyTo,dAmt,{NUM_ONLY},{flt})',
        "E": f'=SUMIFS(dAmt,dPer,">="&cyFrom,dPer,"<="&cyTo,dAmt,{NUM_ONLY},{flt})',
        "F": '=IF(E8=0,"",1)',
        "G": f'=SUMIFS(dQty,dPer,">="&pyFrom,dPer,"<="&pyTo,dQty,{NUM_ONLY},{flt})',
        "H": f'=SUMIFS(dQty,dPer,">="&cyFrom,dPer,"<="&cyTo,dQty,{NUM_ONLY},{flt})',
        "I": '=IF(OR(NOT(mapOK),G8=0),"",D8/G8)',
        "J": '=IF(OR(NOT(mapOK),H8=0),"",E8/H8)',
        "K": '=IF(OR(I8="",J8=""),"",J8-I8)',
        "L": '=IF(OR(I8="",J8=""),"",IF(I8=0,"",J8/I8-1))',
        "M": '=IF(AND(G8=0,H8=0),"수량 없음",IF(G8=0,"전년 수량 없음",IF(H8=0,"당해 수량 없음","")))',
    }
    fmts = {"D": FMT_AMT, "E": FMT_AMT, "F": FMT_PCT, "G": FMT_AMT, "H": FMT_AMT,
            "I": FMT_ASP, "J": FMT_ASP, "K": FMT_DASP, "L": FMT_DPCT}
    top_bottom = Border(top=MED_NAVY, bottom=MED_NAVY, left=THIN, right=THIN)
    style(ws["B8"], F_BOLD, FILL_TOTAL, border=top_bottom)
    for col, value in total.items():
        c = style(ws[f"{col}8"], F_BOLD, FILL_TOTAL, fmt=fmts.get(col),
                  align=LEFT if col in "CM" else None, border=top_bottom)
        c.value = value
    ws.row_dimensions[8].height = 20

    # 업체 행: _calc 시트에서 순위대로 가져오기
    last = CAPACITY + 1

    def pick(calc_col, r):   # _calc 시트의 같은 업체 행에서 값 가져오기
        return f'=IF($N{r}="","",INDEX({Q_CALC}!${calc_col}$2:${calc_col}${last},$N{r}))'

    for i in range(CAPACITY):
        r = 9 + i
        ws[f"B{r}"] = '=IF(ROW()-ROW($B$8)>nComp,"",ROW()-ROW($B$8))'
        ws[f"N{r}"] = f'=IF($B{r}="","",MATCH($B{r},{Q_CALC}!$L$2:$L${last},0))'
        ws[f"C{r}"] = pick("A", r)
        ws[f"D{r}"] = pick("C", r)
        ws[f"E{r}"] = pick("D", r)
        ws[f"F{r}"] = f'=IF($N{r}="","",IF($E$8=0,"",E{r}/$E$8))'
        ws[f"G{r}"] = pick("E", r)
        ws[f"H{r}"] = pick("F", r)
        ws[f"I{r}"] = pick("H", r)
        ws[f"J{r}"] = pick("I", r)
        ws[f"K{r}"] = f'=IF($N{r}="","",IF(OR(I{r}="",J{r}=""),"",J{r}-I{r}))'
        ws[f"L{r}"] = pick("J", r)
        ws[f"M{r}"] = (f'=IF($N{r}="","",IF(AND(G{r}=0,H{r}=0),"수량 없음",IF(G{r}=0,"신규 (전년 수량 없음)",'
                       f'IF(H{r}=0,"당해 수량 없음",""))))')
        for col in "BCDEFGHIJKLMN":
            c = ws[f"{col}{r}"]
            c.font = F_NOTE if col == "M" else F_BASE
            if col in fmts:
                c.number_format = fmts[col]
            if col == "B":
                c.alignment = Alignment(horizontal="center")
    data_rng = f"B9:M{8 + CAPACITY}"
    ws.conditional_formatting.add(data_rng, Rule(
        type="expression", formula=['AND($C9<>"",MOD(ROW(),2)=0)'],
        dxf=DifferentialStyle(fill=PatternFill("solid", bgColor="F7F9FC"))))
    ws.conditional_formatting.add(data_rng, Rule(
        type="expression", formula=['$C9<>""'],
        dxf=DifferentialStyle(border=Border(bottom=Side(style="thin", color="D9D9D9")))))
    group_starts = " ".join(f"{c}9:{c}{8 + CAPACITY}" for c in "DGIM")   # 매출액 | 수량 | ASP | 비고 구분선
    ws.conditional_formatting.add(group_starts, Rule(
        type="expression", formula=['$C9<>""'],
        dxf=DifferentialStyle(border=Border(left=Side(style="thin", color="A6A6A6"),
                                            bottom=Side(style="thin", color="D9D9D9")))))

    ws.freeze_panes = "D9"
    ws.print_title_rows = "6:7"
    # 인쇄 영역: 표시된 업체 행까지만 (빈 수식 행 3,000개를 인쇄하지 않도록)
    ws.defined_names["_xlnm.Print_Area"] = DefinedName(
        "_xlnm.Print_Area", attr_text=f"{Q_OUT}!$A$1:INDEX({Q_OUT}!$M:$M,8+MAX(1,nComp))")
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True


def _text_width(v):
    if v is None:
        return 0
    if isinstance(v, (dt.date, dt.datetime)):
        return 10
    if isinstance(v, (int, float)):
        return len(f"{v:,.0f}")
    return sum(2 if ord(ch) > 0x2E7F else 1 for ch in str(v))


def build_raw(ws, raw_rows):
    for row in raw_rows:
        ws.append(row)
    for c in ws[1]:
        if c.value is not None:
            c.font = F_BOLD
            c.fill = FILL_SUB
    for col_cells in ws.iter_cols(min_row=1, max_row=min(ws.max_row, 300)):
        width = max(_text_width(c.value) for c in col_cells)
        ws.column_dimensions[col_cells[0].column_letter].width = min(max(width + 2, 8), 40)
    for col_cells in ws.iter_cols(min_row=2, max_row=ws.max_row):
        for c in col_cells:
            if isinstance(c.value, (dt.date, dt.datetime)):
                c.number_format = FMT_PERIOD
            elif isinstance(c.value, (int, float)) and abs(c.value) >= 1000 and c.column_letter not in ("A", "J"):
                c.number_format = "#,##0"
    ws.freeze_panes = "A2"


def build_calc(ws, static_companies=None):
    last = CAPACITY + 1
    heads = ["업체", "SUMIFS 조건", "매출 전년", "매출 당해", "수량 전년", "수량 당해", "활성",
             "ASP 전년", "ASP 당해", "ASP 증감률", "정렬 키", "순위"]
    for i, h in enumerate(heads):
        ws.cell(1, 1 + i, h).font = F_BOLD
    flt = "dFlt1,critFlt1,dFlt2,critFlt2,dFlt3,critFlt3"
    if static_companies is None:
        # Excel 2021 / 365: 당해·전년 기간에 거래가 있는 업체를 중복 없이 나열 (아래로 자동 확장)
        formula = ('=IFERROR(_xlfn.UNIQUE(_xlfn._xlws.FILTER(dCust,IFERROR((dCust<>"")*'
                   '((dPer>=cyFrom)*(dPer<=cyTo)+(dPer>=pyFrom)*(dPer<=pyTo)),0))),"")')
        ws["A2"] = ArrayFormula("A2", formula)
    else:
        for i, name in enumerate(static_companies):
            ws.cell(2 + i, 1, name)
    for r in range(2, last + 1):
        ws[f"B{r}"] = (f'=IFERROR(IF(A{r}="","","="&SUBSTITUTE(SUBSTITUTE(SUBSTITUTE(A{r},"~","~~"),'
                       f'"*","~*"),"?","~?")),"")')
        for col, rng_, p0, p1 in (("C", "dAmt", "pyFrom", "pyTo"), ("D", "dAmt", "cyFrom", "cyTo"),
                                  ("E", "dQty", "pyFrom", "pyTo"), ("F", "dQty", "cyFrom", "cyTo")):
            ws[f"{col}{r}"] = (f'=IF($B{r}="","",SUMIFS({rng_},dCust,$B{r},dPer,">="&{p0},dPer,"<="&{p1},'
                               f'{rng_},{NUM_ONLY},{flt}))')
        ws[f"G{r}"] = f'=IF($B{r}="",0,IFERROR(IF(OR(C{r}<>0,D{r}<>0,E{r}<>0,F{r}<>0),1,0),1))'
        ws[f"H{r}"] = f'=IF(OR(G{r}=0,E{r}=0),"",C{r}/E{r})'
        ws[f"I{r}"] = f'=IF(OR(G{r}=0,F{r}=0),"",D{r}/F{r})'
        ws[f"J{r}"] = f'=IF(OR(H{r}="",I{r}=""),"",IF(H{r}=0,"",I{r}/H{r}-1))'
        # 정렬 키: 비활성 업체는 맨 뒤, ASP 증감률을 못 구하는 업체(신규 등)는 활성 업체 중 맨 뒤
        ws[f"K{r}"] = (f'=IF($B{r}="","",IFERROR(IF(G{r}=0,-2E+300,CHOOSE(sortIdx,ROUND(D{r},0),'
                       f'IF(J{r}="",-1E+300,ROUND(-J{r},6)),IF(J{r}="",-1E+300,ROUND(J{r},6)))),-1.5E+300))')
        ws[f"L{r}"] = f'=IF(K{r}="","",RANK(K{r},$K$2:$K${last},0)+COUNTIF($K$2:K{r},K{r})-1)'
    params = [
        ("P2", "표시 업체 수", "Q2", f"=IF(mapOK,SUM($G$2:$G${last}),0)"),
        ("P3", "헤더 마지막 열", "Q3", '=IFERROR(LOOKUP(2,1/(hdrRow<>""),COLUMN(hdrRow)),1)'),
        ("P4", "업체 목록 개수", "Q4", '=COUNTA($A$2:$A$1048576)-IFERROR(IF($A$2="",1,0),0)'),
        ("P5", "필수 열 모두 찾음", "Q5", f'=COUNTIF({Q_SET}!$E$9:$E$12,"✖*")=0'),
    ]
    for lref, label, vref, formula in params:
        ws[lref] = label
        ws[vref] = formula
    # UNIQUE를 못 쓰는 구버전 Excel이면 TRUE (A2의 IFERROR가 오류를 감추므로 따로 확인)
    ws["P6"] = "동적 배열 미지원"
    ws["Q6"] = ArrayFormula("Q6", "=ISERROR(_xlfn.UNIQUE({1;1}))")
    ws["P8"] = "정렬 옵션"
    for i, opt in enumerate(SORT_OPTIONS):
        ws[f"Q{8 + i}"] = opt
    ws.column_dimensions["A"].width = 24
    ws.column_dimensions["P"].width = 16
    ws.column_dimensions["Q"].width = 30
    ws.sheet_state = "hidden"


def define_names(wb):
    s = Q_SET
    add_name(wb, "hdrRowNum", f"{s}!$H$5")
    add_name(wb, "hdrRow", "INDEX(RAW!$A:$XFD,hdrRowNum,0)")
    add_name(wb, "hdrList", f"INDEX(RAW!$A:$XFD,hdrRowNum,1):INDEX(RAW!$A:$XFD,hdrRowNum,{Q_CALC}!$Q$3)")
    cols = {"Per": 9, "Cust": 10, "Qty": 11, "Amt": 12, "Flt1": 22, "Flt2": 23, "Flt3": 24}
    for key, row in cols.items():
        add_name(wb, f"col{key}", f"{s}!$H${row}")
        # 헤더 아래 첫 행부터 시트 끝까지 (열 위치는 헤더명으로 찾은 값)
        add_name(wb, f"d{key}", f"INDEX(RAW!$A:$XFD,hdrRowNum+1,col{key}):INDEX(RAW!$A:$XFD,1048576,col{key})")
    for n, row in (("critFlt1", 22), ("critFlt2", 23), ("critFlt3", 24)):
        add_name(wb, n, f"{s}!$I${row}")
    add_name(wb, "cyFrom", f"{s}!$C$16")
    add_name(wb, "cyTo", f"{s}!$D$16")
    add_name(wb, "pyFrom", f"{s}!$C$18")
    add_name(wb, "pyTo", f"{s}!$D$18")
    add_name(wb, "cyLabel", f"{s}!$E$16")
    add_name(wb, "pyLabel", f"{s}!$E$18")
    add_name(wb, "sortIdx", f"{s}!$H$28")
    add_name(wb, "sortList", f"{Q_CALC}!$Q$8:$Q${7 + len(SORT_OPTIONS)}")
    add_name(wb, "mapOK", f"{Q_CALC}!$Q$5")
    add_name(wb, "dynNA", f"{Q_CALC}!$Q$6")
    add_name(wb, "nComp", f"{Q_CALC}!$Q$2")
    add_name(wb, "nList", f"{Q_CALC}!$Q$4")


def build_workbook(path, raw_rows, settings=None, static_companies=None):
    """raw_rows: RAW 시트에 들어갈 행 목록(헤더 포함).
    static_companies: 검증용. 주면 _calc!A열에 업체 목록을 값으로 넣고, 없으면 UNIQUE 수식 사용."""
    cfg = dict(DEFAULT_SETTINGS)
    cfg.update(settings or {})
    wb = Workbook()
    wb._fonts = IndexedList([Font(name=FONT_NAME, size=10)])   # 기본 글꼴
    ws_guide = wb.active
    ws_guide.title = SH_GUIDE
    ws_out = wb.create_sheet(SH_OUT)
    ws_set = wb.create_sheet(SH_SET)
    ws_raw = wb.create_sheet(SH_RAW)
    ws_calc = wb.create_sheet(SH_CALC)

    build_guide(ws_guide)
    build_settings(ws_set, cfg)
    build_output(ws_out)
    build_raw(ws_raw, raw_rows)
    build_calc(ws_calc, static_companies)
    define_names(wb)

    ws_guide.sheet_properties.tabColor = "808080"
    ws_out.sheet_properties.tabColor = NAVY
    ws_set.sheet_properties.tabColor = "FFC000"
    ws_raw.sheet_properties.tabColor = "70AD47"
    for ws in (ws_guide, ws_set):
        ws.page_setup.orientation = "landscape"
        ws.page_setup.paperSize = ws.PAPERSIZE_A4
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 1
        ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws_set.print_area = "A1:F45"
    wb.active = wb.sheetnames.index(SH_OUT)
    for ws in wb.worksheets:
        ws.sheet_view.tabSelected = ws.title == SH_OUT
    wb.calculation.fullCalcOnLoad = True
    wb.save(path)
    if static_companies is None:
        make_dynamic_array(path, SH_CALC, ["A2", "Q6"])


def sample_raw_rows(df):
    rows = [list(df.columns)]
    for rec in df.itertuples(index=False):
        rows.append([v.to_pydatetime() if isinstance(v, pd.Timestamp) else
                     (v.item() if isinstance(v, np.generic) else v) for v in rec])
    return rows


# ---------------------------------------------------------------------------
# openpyxl은 '동적 배열(자동 확장)' 수식을 직접 쓰지 못하므로, 저장 후 XML에 메타데이터를 추가
# (Excel 2021 / 365가 UNIQUE 같은 수식을 저장할 때 쓰는 형식과 동일)
# ---------------------------------------------------------------------------
_METADATA_XML = (
    '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
    '<metadata xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main" '
    'xmlns:xda="http://schemas.microsoft.com/office/spreadsheetml/2017/dynamicarray">'
    '<metadataTypes count="1"><metadataType name="XLDAPR" minSupportedVersion="120000" copy="1" '
    'pasteAll="1" pasteValues="1" merge="1" splitFirst="1" rowColShift="1" clearFormats="1" '
    'clearComments="1" assign="1" coerce="1" cellMeta="1"/></metadataTypes>'
    '<futureMetadata name="XLDAPR" count="1"><bk><extLst>'
    '<ext uri="{bdbb8cdc-fa1e-496e-a857-3c3f30c029c3}">'
    '<xda:dynamicArrayProperties fDynamic="1" fCollapsed="0"/></ext></extLst></bk></futureMetadata>'
    '<cellMetadata count="1"><bk><rc t="1" v="0"/></bk></cellMetadata></metadata>'
)


def _sheet_xml_path(z, sheet_name):
    wb_xml = z.read("xl/workbook.xml").decode("utf-8")
    rels = z.read("xl/_rels/workbook.xml.rels").decode("utf-8")
    rid = re.search(r'<sheet [^>]*name="%s"[^>]*r:id="([^"]+)"' % re.escape(sheet_name), wb_xml).group(1)
    target = re.search(r'<Relationship [^>]*Id="%s"[^>]*Target="([^"]+)"' % rid, rels)
    if target is None:
        target = re.search(r'<Relationship [^>]*Target="([^"]+)"[^>]*Id="%s"' % rid, rels)
    t = target.group(1).lstrip("/")
    return t if t.startswith("xl/") else "xl/" + t


def make_dynamic_array(path, sheet_name, cells):
    tmp = str(path) + ".tmp"
    with zipfile.ZipFile(path) as z:
        sheet_path = _sheet_xml_path(z, sheet_name)
        with zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as out:
            for item in z.infolist():
                data = z.read(item.filename)
                if item.filename == sheet_path:
                    xml = data.decode("utf-8")
                    for ref in cells:
                        xml, n = re.subn(r'<c r="%s"(?=[ >])' % ref, f'<c r="{ref}" cm="1"', xml, count=1)
                        if n != 1:
                            raise ValueError(f"{sheet_name}!{ref} not found")
                    data = xml.encode("utf-8")
                elif item.filename == "xl/_rels/workbook.xml.rels":
                    data = data.decode("utf-8").replace(
                        "</Relationships>",
                        '<Relationship Id="rIdMeta1" Type="http://schemas.openxmlformats.org/officeDocument/'
                        '2006/relationships/sheetMetadata" Target="metadata.xml"/></Relationships>').encode("utf-8")
                elif item.filename == "[Content_Types].xml":
                    data = data.decode("utf-8").replace(
                        "</Types>",
                        '<Override PartName="/xl/metadata.xml" ContentType="application/'
                        'vnd.openxmlformats-officedocument.spreadsheetml.sheetMetadata+xml"/></Types>').encode("utf-8")
                out.writestr(item, data)
            out.writestr("xl/metadata.xml", _METADATA_XML)
    shutil.move(tmp, path)


def main():
    out = Path(__file__).with_name(OUTPUT_NAME)
    df = make_sample_data()
    build_workbook(out, sample_raw_rows(df))
    print(f"saved: {out}  (예시 데이터 {len(df):,}행)")


if __name__ == "__main__":
    main()
